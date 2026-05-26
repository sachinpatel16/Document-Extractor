import io
import pypdf
import docx
from typing import Optional

def extract_text_from_pdf(file_stream: io.BytesIO) -> str:
    """Extracts text from a PDF file stream using pypdf."""
    try:
        reader = pypdf.PdfReader(file_stream)
        text_parts = []
        for i, page in enumerate(reader.pages):
            text = page.extract_text()
            if text:
                text_parts.append(text)
        return "\n".join(text_parts)
    except Exception as e:
        raise RuntimeError(f"Failed to parse PDF file: {str(e)}")

def extract_text_from_docx(file_bytes: bytes) -> str:
    """Extracts text from a DOCX file using python-docx."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        text_parts = []
        for paragraph in doc.paragraphs:
            if paragraph.text:
                text_parts.append(paragraph.text)
        # Also extract text from tables inside docx
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells if cell.text]
                if row_text:
                    text_parts.append(" | ".join(row_text))
        return "\n".join(text_parts)
    except Exception as e:
        raise RuntimeError(f"Failed to parse DOCX file: {str(e)}")

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extracts text from a TXT file."""
    try:
        return file_bytes.decode("utf-8", errors="ignore")
    except Exception as e:
        raise RuntimeError(f"Failed to parse TXT file: {str(e)}")

def extract_text(uploaded_file) -> str:
    """Detects file type and extracts all text from it."""
    file_bytes = uploaded_file.read()
    filename = uploaded_file.name
    # Reset read pointer for downstream operations if needed
    uploaded_file.seek(0)
    
    if filename.lower().endswith(".pdf"):
        return extract_text_from_pdf(io.BytesIO(file_bytes))
    elif filename.lower().endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    elif filename.lower().endswith(".txt"):
        return extract_text_from_txt(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {filename}. Please upload a .pdf, .docx, or .txt file.")

def format_json_response(json_data: dict) -> str:
    """Helper to convert dictionary to structured representation if needed."""
    import json
    return json.dumps(json_data, indent=2)

def get_custom_css() -> str:
    """Returns custom CSS styles for the Streamlit UI to give a premium look."""
    return """
    <style>
    /* Custom fonts and headers */
    @import url('https://fonts.googleapis.com/css2?family=Outfit:wght@300;400;500;600;700&family=Inter:wght@300;400;500;600;700&display=swap');
    
    :root {
        --primary: #8b5cf6;
        --primary-hover: #7c3aed;
        --secondary: #d946ef;
        --bg-dark: #fafafa;
        --text-main: #1e293b;
        --text-muted: #64748b;
    }
    
    html, body, [class*="css"] {
        font-family: 'Outfit', 'Inter', sans-serif !important;
        font-size: 1rem !important;
    }
    
    /* Scrollbars */
    ::-webkit-scrollbar {
        width: 8px;
        height: 8px;
    }
    ::-webkit-scrollbar-track {
        background: rgba(0, 0, 0, 0.05);
    }
    ::-webkit-scrollbar-thumb {
        background: rgba(0, 0, 0, 0.2);
        border-radius: 4px;
    }
    ::-webkit-scrollbar-thumb:hover {
        background: rgba(0, 0, 0, 0.3);
    }
    
    /* Title text gradient */
    .gradient-text {
        background: linear-gradient(135deg, #a78bfa, #f472b6);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        font-weight: 700;
        font-size: 2.2rem;
        letter-spacing: -0.02em;
    }
    
    /* Navigation styling */
    .sidebar-user {
        padding: 1.5rem;
        background: linear-gradient(135deg, rgba(139, 92, 246, 0.08), rgba(217, 70, 239, 0.04));
        border-radius: 16px;
        margin-bottom: 1.5rem;
        border: 1px solid rgba(139, 92, 246, 0.15);
        box-shadow: 0 4px 20px rgba(139, 92, 246, 0.08);
    }
    
    /* Style radio group for navigation */
    div[data-testid="stRadio"] > label {
        font-size: 0.85rem !important;
        font-weight: 600 !important;
        color: var(--text-muted) !important;
        text-transform: uppercase;
        letter-spacing: 0.05em;
        margin-bottom: 0.5rem;
    }
    
    div[data-testid="stRadio"] div[role="radiogroup"] {
        gap: 0.5rem;
        background: rgba(255, 255, 255, 0.9) !important;
        border: 1px solid rgba(139, 92, 246, 0.1) !important;
        border-radius: 12px;
        padding: 0.6rem !important;
        box-shadow: 0 2px 8px rgba(0, 0, 0, 0.05) !important;
    }
    
    div[data-testid="stRadio"] div[role="radiogroup"] label {
        padding: 0.75rem 1rem !important;
        border-radius: 8px !important;
        background-color: transparent !important;
        border: 1px solid transparent !important;
        transition: all 0.2s ease !important;
        color: #475569 !important;
        font-size: 0.95rem !important;
        font-weight: 500 !important;
    }
    
    div[data-testid="stRadio"] div[role="radiogroup"] label:hover {
        background-color: rgba(139, 92, 246, 0.08) !important;
        border-color: rgba(139, 92, 246, 0.15) !important;
    }
    
    div[data-testid="stRadio"] div[role="radiogroup"] label[data-checked="true"] {
        background: linear-gradient(135deg, rgba(139, 92, 246, 0.12), rgba(217, 70, 239, 0.06)) !important;
        border: 1px solid rgba(139, 92, 246, 0.2) !important;
        color: #8b5cf6 !important;
        font-weight: 600 !important;
        box-shadow: 0 2px 8px rgba(139, 92, 246, 0.15) !important;
    }
    
    /* Sidebar divider */
    [data-testid="stSidebar"] > div:first-child > div > div > div > hr {
        border-color: rgba(139, 92, 246, 0.1) !important;
        margin: 1.5rem 0 !important;
    }
    
    /* Custom container borders styling (cards) */
    div[data-testid="stVerticalBlockBorderWrapper"] {
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.95), rgba(248, 250, 252, 0.9)) !important;
        border: 1px solid rgba(139, 92, 246, 0.08) !important;
        border-radius: 16px !important;
        padding: 1.5rem !important;
        transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1) !important;
        margin-bottom: 1.25rem !important;
        backdrop-filter: blur(8px) !important;
        font-size: 1rem !important;
    }
    div[data-testid="stVerticalBlockBorderWrapper"]:hover {
        border-color: rgba(139, 92, 246, 0.15) !important;
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.98), rgba(248, 250, 252, 0.95)) !important;
        box-shadow: 0 8px 30px rgba(139, 92, 246, 0.08) !important;
    }
    
    /* File Uploader styling */
    [data-testid="stFileUploader"] {
        background: linear-gradient(135deg, rgba(248, 250, 252, 0.9), rgba(241, 245, 249, 0.85)) !important;
        border: 1px solid rgba(139, 92, 246, 0.1) !important;
        border-radius: 16px !important;
        padding: 1.75rem !important;
        transition: all 0.3s ease !important;
        font-size: 1rem !important;
    }
    [data-testid="stFileUploader"]:hover {
        border-color: rgba(139, 92, 246, 0.2) !important;
        background: linear-gradient(135deg, rgba(248, 250, 252, 0.95), rgba(241, 245, 249, 0.9)) !important;
    }
    
    /* Forms styling */
    [data-testid="stForm"] {
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.95), rgba(248, 250, 252, 0.9)) !important;
        border: 1px solid rgba(139, 92, 246, 0.08) !important;
        border-radius: 16px !important;
        padding: 1.75rem !important;
        box-shadow: 0 8px 30px rgba(139, 92, 246, 0.06) !important;
        backdrop-filter: blur(12px) !important;
        font-size: 1rem !important;
    }
    
    /* Input fields overrides */
    div[data-baseweb="input"] {
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.9), rgba(248, 250, 252, 0.85)) !important;
        border: 1px solid rgba(139, 92, 246, 0.08) !important;
        border-radius: 8px !important;
        transition: all 0.3s ease !important;
        font-size: 1rem !important;
    }
    div[data-baseweb="input"]:focus-within {
        border-color: rgba(139, 92, 246, 0.2) !important;
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.95), rgba(248, 250, 252, 0.9)) !important;
        box-shadow: 0 0 10px rgba(139, 92, 246, 0.1) !important;
    }
    
    /* Tabs styling */
    button[data-baseweb="tab"] {
        font-size: 0.95rem !important;
        font-weight: 600 !important;
        color: var(--text-muted) !important;
        border-bottom: none !important;
        transition: all 0.3s ease !important;
        padding: 0.6rem 1.5rem !important;
    }
    button[data-baseweb="tab"][aria-selected="true"] {
        color: #c084fc !important;
        border-bottom: none !important;
        font-weight: 700 !important;
    }
    
    /* Primary buttons styling */
    button[data-testid="baseButton-primary"] {
        background: linear-gradient(135deg, #8b5cf6, #d946ef) !important;
        border: none !important;
        color: white !important;
        font-weight: 600 !important;
        border-radius: 12px !important;
        padding: 0.6rem 1.8rem !important;
        box-shadow: 0 4px 15px rgba(139, 92, 246, 0.35) !important;
        transition: all 0.3s cubic-bezier(0.4, 0, 0.2, 1) !important;
    }
    button[data-testid="baseButton-primary"]:hover {
        transform: translateY(-2px) !important;
        box-shadow: 0 8px 25px rgba(139, 92, 246, 0.5) !important;
        background: linear-gradient(135deg, #7c3aed, #c084fc) !important;
    }
    
    /* Secondary buttons styling */
    button[data-testid="baseButton-secondary"] {
        background-color: rgba(255, 255, 255, 0.03) !important;
        border: none !important;
        color: #cbd5e1 !important;
        border-radius: 12px !important;
        font-weight: 500 !important;
        transition: all 0.2s ease !important;
    }
    button[data-testid="baseButton-secondary"]:hover {
        background-color: rgba(255, 255, 255, 0.08) !important;
        border-color: transparent !important;
        color: #ffffff !important;
    }
    
    /* Popover UI adjustments */
    div[data-testid="stPopoverBody"] {
        background-color: #ffffff !important;
        border: none !important;
        border-radius: 16px !important;
        box-shadow: 0 10px 40px rgba(0, 0, 0, 0.15) !important;
    }
    
    /* Document Preview styling */
    .document-preview-container {
        background-color: #f8fafc;
        border: none;
        border-radius: 12px;
        padding: 1.25rem;
        max-height: 350px;
        overflow-y: auto;
    }
    .document-preview-text {
        margin: 0;
        white-space: pre-wrap;
        word-wrap: break-word;
        font-family: 'SFMono-Regular', Consolas, 'Liberation Mono', Menlo, monospace;
        font-size: 0.85rem;
        color: #475569;
        line-height: 1.6;
    }
    
    /* Empty state styling */
    .empty-state-container {
        text-align: center;
        margin-top: 3rem;
        border: none;
        border-radius: 20px;
        padding: 4.5rem 2rem;
        background: rgba(248, 250, 252, 0.8);
        transition: all 0.3s ease;
    }
    .empty-state-container:hover {
        border-color: transparent;
        background: rgba(139, 92, 246, 0.05);
    }
    .empty-state-icon {
        font-size: 3.8rem;
        margin-bottom: 1.25rem;
        animation: float 4s ease-in-out infinite;
    }
    @keyframes float {
        0% { transform: translateY(0px); }
        50% { transform: translateY(-10px); }
        100% { transform: translateY(0px); }
    }
    
    /* Custom loader */
    .extraction-loader {
        text-align: center;
        padding: 4.5rem 2rem;
        background: rgba(255, 255, 255, 0.9);
        border: none;
        border-radius: 20px;
        margin: 2rem 0;
        box-shadow: 0 12px 40px rgba(0, 0, 0, 0.1);
        backdrop-filter: blur(10px);
    }
    .loader-spinner {
        width: 55px;
        height: 55px;
        border: 4px solid rgba(0, 0, 0, 0.1);
        border-top-color: #d946ef;
        border-radius: 50%;
        margin: 0 auto 1.5rem auto;
        animation: spin 1.2s linear infinite;
        box-shadow: 0 0 15px rgba(217, 70, 239, 0.2);
    }
    .loader-title {
        margin: 0 0 0.5rem 0;
        color: #1e293b;
        font-weight: 700;
        font-size: 1.25rem;
        letter-spacing: -0.01em;
    }
    .loader-desc {
        margin: 0;
        color: #64748b;
        font-size: 0.92rem;
        max-width: 80%;
        margin: 0 auto;
        line-height: 1.5;
    }
    @keyframes spin {
        to { transform: rotate(360deg); }
    }
    
    /* Metric Cards */
    .metric-card {
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.95), rgba(248, 250, 252, 0.9));
        border: 1px solid rgba(139, 92, 246, 0.1);
        border-radius: 16px;
        padding: 1.5rem;
        text-align: center;
        box-shadow: 0 4px 20px rgba(139, 92, 246, 0.06);
        transition: all 0.3s ease;
        margin-bottom: 1rem;
        font-size: 1rem;
    }
    .metric-card:hover {
        transform: translateY(-2px);
        border-color: rgba(139, 92, 246, 0.2);
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.98), rgba(248, 250, 252, 0.95));
        box-shadow: 0 8px 25px rgba(139, 92, 246, 0.1);
    }
    .metric-icon {
        font-size: 2rem;
        margin-bottom: 0.6rem;
    }
    .metric-value {
        font-size: 2.2rem;
        font-weight: 700;
        background: linear-gradient(135deg, #a855f7, #ec4899);
        -webkit-background-clip: text;
        -webkit-text-fill-color: transparent;
        margin-bottom: 0.3rem;
    }
    .metric-label {
        font-size: 0.8rem;
        color: #94a3b8;
        text-transform: uppercase;
        letter-spacing: 0.08em;
        font-weight: 600;
    }
    
    /* Smart Report Styling */
    .extracted-report-container {
        display: flex;
        flex-direction: column;
        gap: 1.5rem;
        margin-top: 1rem;
    }
    .report-header-card {
        background: linear-gradient(135deg, rgba(255, 255, 255, 0.95), rgba(248, 250, 252, 0.9));
        border: 1px solid rgba(139, 92, 246, 0.1);
        border-radius: 16px;
        padding: 1.5rem;
        display: flex;
        align-items: center;
        gap: 1.25rem;
        box-shadow: 0 4px 20px rgba(139, 92, 246, 0.06);
        font-size: 1rem;
    }
    .report-header-icon {
        font-size: 2.25rem;
        background: linear-gradient(135deg, rgba(139, 92, 246, 0.1), rgba(217, 70, 239, 0.05));
        width: 60px;
        height: 60px;
        border-radius: 12px;
        display: flex;
        align-items: center;
        justify-content: center;
        border: 1px solid rgba(139, 92, 246, 0.15);
    }
    .report-header-label {
        font-size: 0.75rem;
        color: #a78bfa;
        text-transform: uppercase;
        font-weight: 600;
        letter-spacing: 0.05em;
    }
    .report-header-title {
        font-size: 1.4rem;
        color: #1e293b;
        font-weight: 700;
        margin: 0.25rem 0 0 0 !important;
    }
    .report-section-title {
        font-size: 0.95rem;
        font-weight: 700;
        color: #1e293b;
        text-transform: uppercase;
        letter-spacing: 0.05em;
        margin-top: 0.5rem;
        border-left: 3px solid #8b5cf6;
        padding-left: 0.75rem;
    }
    .contact-grid-card {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(220px, 1fr));
        gap: 1rem;
        background: rgba(248, 250, 252, 0.8);
        border: 1px solid rgba(0, 0, 0, 0.04);
        border-radius: 16px;
        padding: 1.25rem;
    }
    .contact-grid-item {
        display: flex;
        align-items: center;
        gap: 0.85rem;
    }
    .contact-item-icon-wrapper {
        background: rgba(139, 92, 246, 0.1);
        color: #8b5cf6;
        width: 36px;
        height: 36px;
        border-radius: 10px;
        display: flex;
        align-items: center;
        justify-content: center;
        border: 1px solid rgba(139, 92, 246, 0.2);
        flex-shrink: 0;
    }
    .contact-icon {
        width: 18px;
        height: 18px;
    }
    .contact-item-details {
        display: flex;
        flex-direction: column;
        min-width: 0;
    }
    .contact-item-label {
        font-size: 0.7rem;
        color: #64748b;
        text-transform: uppercase;
        font-weight: 600;
    }
    .contact-item-value {
        font-size: 0.85rem;
        color: #1e293b;
        font-weight: 500;
        white-space: nowrap;
        overflow: hidden;
        text-overflow: ellipsis;
    }
    .contact-link {
        color: #a78bfa;
        text-decoration: none;
        transition: color 0.2s ease;
    }
    .contact-link:hover {
        color: #c084fc;
        text-decoration: underline;
    }
    .report-text-block-card {
        background: rgba(248, 250, 252, 0.8);
        border: 1px solid rgba(0, 0, 0, 0.06);
        border-radius: 14px;
        padding: 1.25rem;
    }
    .report-text-paragraph {
        font-size: 0.92rem;
        color: #475569;
        line-height: 1.6;
        margin: 0;
    }
    .flat-fields-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
        gap: 0.85rem;
    }
    .flat-field-card {
        background: rgba(248, 250, 252, 0.8);
        border: 1px solid rgba(0, 0, 0, 0.05);
        border-radius: 12px;
        padding: 1rem;
        display: flex;
        flex-direction: column;
        gap: 0.35rem;
    }
    .flat-field-label {
        font-size: 0.7rem;
        color: #64748b;
        text-transform: uppercase;
        font-weight: 600;
    }
    .flat-field-value {
        font-size: 0.95rem;
        color: #1e293b;
        font-weight: 600;
    }
    .report-badge-container {
        display: flex;
        flex-wrap: wrap;
        gap: 0.5rem;
    }
    .report-badge-pill {
        background: rgba(139, 92, 246, 0.12);
        border: 1px solid rgba(139, 92, 246, 0.25);
        color: #8b5cf6;
        padding: 0.35rem 0.85rem;
        border-radius: 20px;
        font-size: 0.82rem;
        font-weight: 500;
        box-shadow: 0 2px 8px rgba(139, 92, 246, 0.08);
        transition: all 0.25s ease;
    }
    .report-badge-pill:hover {
        background: rgba(139, 92, 246, 0.15);
        border-color: rgba(139, 92, 246, 0.35);
        transform: translateY(-1px);
        box-shadow: 0 4px 12px rgba(139, 92, 246, 0.08);
    }
    .report-timeline {
        display: flex;
        flex-direction: column;
        gap: 1.25rem;
        position: relative;
        padding-left: 1.5rem;
        margin-left: 0.5rem;
        border-left: 2px solid rgba(139, 92, 246, 0.15);
    }
    .timeline-item-card {
        position: relative;
    }
    .timeline-node-dot {
        position: absolute;
        left: calc(-1.5rem - 6px);
        top: 6px;
        width: 10px;
        height: 10px;
        background: #8b5cf6;
        border-radius: 50%;
        border: 2px solid #ffffff;
        box-shadow: 0 0 8px #8b5cf6;
    }
    .timeline-card-content {
        background: rgba(248, 250, 252, 0.8);
        border: 1px solid rgba(0, 0, 0, 0.04);
        border-radius: 14px;
        padding: 1.15rem;
        transition: all 0.25s ease;
    }
    .timeline-card-content:hover {
        background: rgba(255, 255, 255, 0.95);
        border-color: rgba(139, 92, 246, 0.2);
    }
    .timeline-card-title {
        font-size: 1rem;
        font-weight: 600;
        color: #1e293b;
        margin: 0 !important;
    }
    .timeline-card-subtitle {
        font-size: 0.82rem;
        color: #64748b;
        margin-top: 0.35rem;
    }
    .timeline-card-desc {
        font-size: 0.88rem;
        color: #475569;
        line-height: 1.5;
        margin: 0.75rem 0 0 0 !important;
    }
    .timeline-card-extras {
        display: flex;
        flex-wrap: wrap;
        gap: 0.5rem;
        margin-top: 0.75rem;
        padding-top: 0.75rem;
        border-top: 1px dashed rgba(0, 0, 0, 0.05);
    }
    .timeline-card-extra-item {
        font-size: 0.78rem;
        color: #64748b;
        background: rgba(248, 250, 252, 0.8);
        padding: 0.2rem 0.5rem;
        border-radius: 6px;
        border: 1px solid rgba(0, 0, 0, 0.04);
    }
    .dict-fields-grid {
        display: grid;
        grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
        gap: 0.85rem;
    }
    </style>
    """

def render_smart_extracted_data(result: dict) -> str:
    """
    Generates a highly polished, responsive HTML report representing the extracted JSON data.
    Automatically formats structures like lists, objects, contact information, and long paragraphs
    into specialized components (timelines, tag clouds, contact cards, grids).
    """
    import html
    import json
    
    if not isinstance(result, dict):
        return f"<div style='color: #ef4444; padding: 1rem;'>Invalid extraction data format.</div>"
        
    html_out = []
    
    # 1. Look for a primary header field
    header_keys = ['name', 'full_name', 'title', 'document_title', 'invoice_number', 'invoice_no']
    header_val = None
    header_key_found = None
    
    for hk in header_keys:
        for k in result.keys():
            if k.lower() == hk or k.lower().replace('_', '') == hk:
                header_val = result[k]
                header_key_found = k
                break
        if header_val:
            break
            
    # 2. Extract nested sections vs flat fields
    flat_fields = []
    text_blocks = []
    list_fields = []
    dict_fields = []
    contact_field = None
    
    for k, v in result.items():
        if k == header_key_found:
            continue
            
        # Check for contact field
        if k.lower() in ['contact', 'contact_info', 'contact_details', 'contacts'] and isinstance(v, dict):
            contact_field = (k, v)
            continue
            
        if isinstance(v, dict):
            dict_fields.append((k, v))
        elif isinstance(v, list):
            list_fields.append((k, v))
        elif isinstance(v, (str, int, float)) or v is None:
            v_str = str(v) if v is not None else ""
            if len(v_str) > 100 or '\\n' in v_str or '\n' in v_str or k.lower() in ['profile', 'summary', 'description', 'objective', 'about']:
                text_blocks.append((k, v_str))
            else:
                flat_fields.append((k, v_str))
                
    # Build HTML
    html_out.append("<div class='extracted-report-container'>")
    
    # Add a header if found
    if header_val:
        label = header_key_found.replace('_', ' ').title()
        html_out.append(f"""
            <div class='report-header-card'>
                <div class='report-header-icon'>📄</div>
                <div>
                    <span class='report-header-label'>{label}</span>
                    <h2 class='report-header-title'>{html.escape(str(header_val))}</h2>
                </div>
            </div>
        """)
        
    # Render Contact Card if present
    if contact_field:
        k, v = contact_field
        html_out.append(f"""
            <div class='report-section-title'>📍 {k.replace('_', ' ').title()}</div>
            <div class='contact-grid-card'>
        """)
        
        # Parse standard contact fields
        icons = {
            'email': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><path d="M4 4h16c1.1 0 2 .9 2 2v12c0 1.1-.9 2-2 2H4c-1.1 0-2-.9-2-2V6c0-1.1.9-2 2-2z"></path><polyline points="22,6 12,13 2,6"></polyline></svg>',
            'phone': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><path d="M22 16.92v3a2 2 0 0 1-2.18 2 19.79 19.79 0 0 1-8.63-3.07 19.5 19.5 0 0 1-6-6 19.79 19.79 0 0 1-3.07-8.67A2 2 0 0 1 4.11 2h3a2 2 0 0 1 2 1.72 12.84 12.84 0 0 0 .7 2.81 2 2 0 0 1-.45 2.11L8.09 9.91a16 16 0 0 0 6 6l1.27-1.27a2 2 0 0 1 2.11-.45 12.84 12.84 0 0 0 2.81.7A2 2 0 0 1 22 16.92z"></path></svg>',
            'location': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><path d="M12 2a8 8 0 0 0-8 8c0 5.25 8 12 8 12s8-6.75 8-12a8 8 0 0 0-8-8z"></path><circle cx="12" cy="10" r="3"></circle></svg>',
            'linkedin': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><path d="M16 8a6 6 0 0 1 6 6v7h-4v-7a2 2 0 0 0-2-2 2 2 0 0 0-2 2v7h-4v-7a6 6 0 0 1 6-6z"></path><rect x="2" y="9" width="4" height="12"></rect><circle cx="4" cy="4" r="2"></circle></svg>',
            'github': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><path d="M9 19c-5 1.5-5-2.5-7-3m14 6v-3.87a3.37 3.37 0 0 0-.94-2.61c3.14-.35 6.44-1.54 6.44-7A5.44 5.44 0 0 0 20 4.77 5.07 5.07 0 0 0 19.91 1S18.73.65 16 2.48a13.38 13.38 0 0 0-7 0C6.27.65 5.09 1 5.09 1A5.07 5.07 0 0 0 5 4.77a5.44 5.44 0 0 0-1.5 3.78c0 5.42 3.3 6.61 6.44 7A3.37 3.37 0 0 0 9 18.13V22"></path></svg>',
            'website': '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><circle cx="12" cy="12" r="10"></circle><line x1="2" y1="12" x2="22" y2="12"></line><path d="M12 2a15.3 15.3 0 0 1 4 10 15.3 15.3 0 0 1-4 10 15.3 15.3 0 0 1-4-10 15.3 15.3 0 0 1 4-10z"></path></svg>'
        }
        
        default_icon = '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="currentColor" stroke-width="2" stroke-linecap="round" stroke-linejoin="round" class="contact-icon"><circle cx="12" cy="12" r="10"></circle><line x1="12" y1="16" x2="12" y2="12"></line><line x1="12" y1="8" x2="12.01" y2="8"></line></svg>'
        
        for key, val in v.items():
            k_lower = key.lower()
            icon = default_icon
            for ik in icons.keys():
                if ik in k_lower:
                    icon = icons[ik]
                    break
                    
            val_escaped = html.escape(str(val))
            val_html = val_escaped
            if "@" in val_escaped and "." in val_escaped and " " not in val_escaped:
                val_html = f"<a href='mailto:{val_escaped}' class='contact-link'>{val_escaped}</a>"
            elif val_escaped.startswith(('http://', 'https://', 'www.')):
                url = val_escaped if val_escaped.startswith(('http://', 'https://')) else f"https://{val_escaped}"
                val_html = f"<a href='{url}' target='_blank' class='contact-link'>{val_escaped}</a>"
                
            html_out.append(f"""
                <div class='contact-grid-item'>
                    <div class='contact-item-icon-wrapper'>{icon}</div>
                    <div class='contact-item-details'>
                        <span class='contact-item-label'>{key.replace('_', ' ').title()}</span>
                        <span class='contact-item-value'>{val_html}</span>
                    </div>
                </div>
            """)
        html_out.append("</div>")
        
    # Render Text Blocks
    if text_blocks:
        for k, v in text_blocks:
            html_out.append(f"""
                <div class='report-section-title'>📝 {k.replace('_', ' ').title()}</div>
                <div class='report-text-block-card'>
                    <p class='report-text-paragraph'>{html.escape(v).replace('\\n', '<br>').replace('\n', '<br>')}</p>
                </div>
            """)
            
    # Render Flat Fields
    if flat_fields:
        html_out.append("""
            <div class='report-section-title'>📊 General Details</div>
            <div class='flat-fields-grid'>
        """)
        for k, v in flat_fields:
            html_out.append(f"""
                <div class='flat-field-card'>
                    <span class='flat-field-label'>{k.replace('_', ' ').title()}</span>
                    <span class='flat-field-value'>{html.escape(v)}</span>
                </div>
            """)
        html_out.append("</div>")
        
    # Render List Fields
    if list_fields:
        for k, v in list_fields:
            if not v:
                continue
                
            # List of simple scalars -> Pills
            if all(isinstance(item, (str, int, float)) for item in v):
                html_out.append(f"""
                    <div class='report-section-title'>⚡ {k.replace('_', ' ').title()}</div>
                    <div class='report-badge-container'>
                """)
                for item in v:
                    html_out.append(f"<span class='report-badge-pill'>{html.escape(str(item))}</span>")
                html_out.append("</div>")
                
            # List of dicts -> Timeline
            elif all(isinstance(item, dict) for item in v):
                html_out.append(f"""
                    <div class='report-section-title'>📜 {k.replace('_', ' ').title()}</div>
                    <div class='report-timeline'>
                """)
                
                for idx, item in enumerate(v):
                    title_keys = ['title', 'role', 'designation', 'position', 'degree', 'course', 'item_name', 'name']
                    subtitle_keys = ['company', 'employer', 'school', 'institution', 'university', 'vendor', 'organization']
                    date_keys = ['duration', 'dates', 'period', 'year', 'date', 'time', 'start_date', 'end_date']
                    desc_keys = ['description', 'summary', 'details', 'responsibilities', 'highlights', 'bullet_points']
                    
                    title_val = ""
                    subtitle_val = ""
                    date_val = ""
                    desc_val = ""
                    extra_vals = []
                    
                    for sub_k, sub_v in item.items():
                        sub_k_lower = sub_k.lower()
                        if any(tk in sub_k_lower for tk in title_keys) and not title_val:
                            title_val = str(sub_v)
                        elif any(sk in sub_k_lower for sk in subtitle_keys) and not subtitle_val:
                            subtitle_val = str(sub_v)
                        elif any(dk in sub_k_lower for dk in date_keys) and not date_val:
                            date_val = str(sub_v)
                        elif any(dk in sub_k_lower for dk in desc_keys) and not desc_val:
                            if isinstance(sub_v, list):
                                desc_val = "\\n".join([f"• {x}" for x in sub_v])
                            else:
                                desc_val = str(sub_v)
                        else:
                            if isinstance(sub_v, list):
                                extra_vals.append(f"<strong>{sub_k.replace('_', ' ').title()}:</strong> " + ", ".join([str(x) for x in sub_v]))
                            elif isinstance(sub_v, dict):
                                extra_vals.append(f"<strong>{sub_k.replace('_', ' ').title()}:</strong> " + json.dumps(sub_v))
                            else:
                                extra_vals.append(f"<strong>{sub_k.replace('_', ' ').title()}:</strong> {sub_v}")
                                
                    title_html = title_val or f"Record {idx + 1}"
                    subtitle_html = subtitle_val
                    if date_val:
                        subtitle_html = f"{subtitle_html} &nbsp;|&nbsp; 🕰️ {date_val}" if subtitle_html else date_val
                        
                    html_out.append(f"""
                        <div class='timeline-item-card'>
                            <div class='timeline-node-dot'></div>
                            <div class='timeline-card-content'>
                                <h4 class='timeline-card-title'>{html.escape(title_html)}</h4>
                                {f"<div class='timeline-card-subtitle'>{html.escape(subtitle_html)}</div>" if subtitle_html else ""}
                                {f"<p class='timeline-card-desc'>{html.escape(desc_val).replace('\\\\n', '<br>').replace('\\n', '<br>').replace('\n', '<br>')}</p>" if desc_val else ""}
                    """)
                    
                    if extra_vals:
                        html_out.append("<div class='timeline-card-extras'>")
                        for ev in extra_vals:
                            html_out.append(f"<div class='timeline-card-extra-item'>{ev}</div>")
                        html_out.append("</div>")
                        
                    html_out.append("""
                            </div>
                        </div>
                    """)
                html_out.append("</div>")
                
    # Render generic dicts
    if dict_fields:
        for k, v in dict_fields:
            html_out.append(f"""
                <div class='report-section-title'>🗂️ {k.replace('_', ' ').title()}</div>
                <div class='dict-fields-grid'>
            """)
            for sub_k, sub_v in v.items():
                if isinstance(sub_v, (dict, list)):
                    sub_v_str = json.dumps(sub_v)
                else:
                    sub_v_str = str(sub_v)
                html_out.append(f"""
                    <div class='flat-field-card'>
                        <span class='flat-field-label'>{sub_k.replace('_', ' ').title()}</span>
                        <span class='flat-field-value'>{html.escape(sub_v_str)}</span>
                    </div>
                """)
            html_out.append("</div>")
            
    html_out.append("</div>")
    return "\\n".join(html_out)

